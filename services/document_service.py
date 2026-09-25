import io
import csv

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def extract_text(
    data,
    filename,
    mime_type
):

    filename_lower = (
        filename or ""
    ).lower()

    # ========================================================
    # TEXT
    # ========================================================

    if (
        mime_type.startswith("text/")
        or filename_lower.endswith(
            (".txt", ".csv", ".md", ".log")
        )
    ):

        return data.decode(
            "utf-8",
            errors="ignore"
        )

    # ========================================================
    # PDF
    # ========================================================

    if (
        mime_type == "application/pdf"
        or filename_lower.endswith(".pdf")
    ):

        reader = PdfReader(
            io.BytesIO(data)
        )

        pages = []

        for page in reader.pages:

            text = page.extract_text()

            if text:

                pages.append(text)

        return "\n\n".join(pages)

    # ========================================================
    # DOCX
    # ========================================================

    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or filename_lower.endswith(".docx")
    ):

        document = Document(
            io.BytesIO(data)
        )

        paragraphs = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                paragraphs.append(
                    paragraph.text
                )

        return "\n".join(
            paragraphs
        )

    # ========================================================
    # XLSX
    # ========================================================

    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        or filename_lower.endswith(".xlsx")
    ):

        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True
        )

        lines = []

        for worksheet in workbook.worksheets:

            lines.append(
                f"--- Sheet: {worksheet.title} ---"
            )

            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:

                    lines.append(
                        " | ".join(values)
                    )

        return "\n".join(lines)

    # ========================================================
    # PPTX
    # ========================================================

    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        or filename_lower.endswith(".pptx")
    ):

        presentation = Presentation(
            io.BytesIO(data)
        )

        lines = []

        for index, slide in enumerate(
            presentation.slides,
            start=1
        ):

            lines.append(
                f"--- Slide {index} ---"
            )

            for shape in slide.shapes:

                if hasattr(
                    shape,
                    "text"
                ):

                    text = shape.text.strip()

                    if text:

                        lines.append(
                            text
                        )

        return "\n".join(lines)

    # ========================================================
    # Google CSV / OTHER
    # ========================================================

    try:

        return data.decode(
            "utf-8",
            errors="ignore"
        )

    except Exception:

        raise RuntimeError(
            f"Unsupported document type: "
            f"{mime_type}"
        )