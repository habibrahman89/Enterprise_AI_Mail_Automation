import os
import re
import logging
from pathlib import Path

log = logging.getLogger(__name__)

BASE = Path(os.getenv('ATTACHMENT_STORAGE', 'storage/attachments'))

MAX_ATTACHMENT_BYTES = int(os.getenv('ATTACHMENT_MAX_SIZE_MB', '25')) * 1024 * 1024

# Extraction is only attempted for known-safe document types; anything
# else is stored but left unextracted rather than executed/opened.
EXTRACTABLE_EXTENSIONS = {
    '.pdf', '.docx', '.xlsx', '.xlsm', '.xltx', '.xltm',
    '.txt', '.csv', '.log', '.png', '.jpg', '.jpeg', '.tiff', '.bmp',
}


def safe_filename(n):
    return re.sub(r'[^A-Za-z0-9._ -]', '_', os.path.basename(n or 'attachment'))[:180]


def save_bytes(email_id, filename, data):
    if data and len(data) > MAX_ATTACHMENT_BYTES:
        log.warning(
            "Rejected attachment %r for email %s: %d bytes exceeds limit of %d bytes",
            filename, email_id, len(data), MAX_ATTACHMENT_BYTES,
        )
        raise ValueError(
            f"Attachment '{filename}' exceeds the maximum allowed size "
            f"({MAX_ATTACHMENT_BYTES // (1024 * 1024)} MB) and was not stored."
        )

    d = BASE / str(email_id)
    d.mkdir(parents=True, exist_ok=True)
    p = d / safe_filename(filename)
    p.write_bytes(data)
    return str(p)


def extract_attachment_text(path, content_type=''):
    if not path or not os.path.exists(path):
        return ''

    ext = Path(path).suffix.lower()
    if ext not in EXTRACTABLE_EXTENSIONS:
        return ''

    try:
        if ext == '.pdf':
            from pypdf import PdfReader
            return '\n'.join(
                (p.extract_text() or '') for p in PdfReader(path).pages
            )[:50000]

        if ext == '.docx':
            from docx import Document
            return '\n'.join(p.text for p in Document(path).paragraphs)[:50000]

        if ext in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append('[SHEET: ' + ws.title + ']')
                for r in ws.iter_rows(values_only=True):
                    out.append(' | '.join('' if v is None else str(v) for v in r))
            return '\n'.join(out)[:50000]

        if ext in ('.txt', '.csv', '.log'):
            return Path(path).read_text(encoding='utf-8', errors='replace')[:50000]

        if ext in ('.png', '.jpg', '.jpeg', '.tiff', '.bmp') and os.getenv('OCR_ENABLED', '0') == '1':
            import pytesseract
            from PIL import Image
            return pytesseract.image_to_string(Image.open(path))[:50000]

    except Exception as exc:
        log.warning("Attachment text extraction failed for %s: %s", path, exc)
        return f'[Extraction failed: {exc}]'

    return ''
